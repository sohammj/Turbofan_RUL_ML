"""Render the reviewed group-activity report source as an editable Word document."""

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "deliverables" / "GROUP_ACTIVITY_REPORT_SOURCE.md"
OUTPUT = ROOT / "deliverables" / "GROUP_ACTIVITY_REPORT_DRAFT.docx"


def clean(value):
    return re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", value).replace("`", "").replace("**", "").strip()


def shade(cell, color):
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(element)


def add_table(doc, rows):
    count = len(rows[0])
    table = doc.add_table(rows=0, cols=count)
    table.autofit = False
    for row_number, values in enumerate(rows):
        cells = table.add_row().cells
        if row_number == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            table.rows[-1]._tr.get_or_add_trPr().append(repeat)
        for col_number, value in enumerate(values):
            cell = cells[col_number]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = clean(value)
            if row_number == 0:
                shade(cell, "233B55")
            elif row_number % 2 == 0:
                shade(cell, "F0F3F6")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.line_spacing = 1
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9.5 if count > 4 else 10)
                    if row_number == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_pipeline(doc):
    stages = [
        "NASA source files", "Quality checks", "EDA and features", "Engine-wise validation",
        "Model selection", "Refit on training", "Official test", "Error and stability audit",
    ]
    table = doc.add_table(rows=2, cols=4)
    table.autofit = False
    for index, stage in enumerate(stages):
        cell = table.cell(index // 4, index % 4)
        cell.text = f"{index + 1}. {stage}"
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, "E9F0F7" if index % 2 == 0 else "F6F8FA")
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(0.8)
    section.top_margin, section.bottom_margin = Inches(0.75), Inches(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1
    normal.paragraph_format.space_after = Pt(6)
    for name, size, before, after in (("Title", 18, 0, 14), ("Heading 1", 14, 13, 6), ("Heading 2", 12, 10, 4)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 1
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    cap = doc.styles["Caption"]
    cap.font.name = "Times New Roman"
    cap.font.size = Pt(10)
    cap.font.italic = True
    cap.font.color.rgb = RGBColor(0, 0, 0)
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(7)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Machine Learning Group Activity Report  |  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(9)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    cover = True
    pending_figure = None
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(clean(line[2:]), style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(90)
        elif line.startswith("## "):
            if cover:
                doc.add_page_break()
                cover = False
            if line == "## Appendix A Repository and reproducibility":
                doc.add_page_break()
            doc.add_paragraph(clean(line[3:]), style="Heading 1")
        elif line.startswith("### "):
            doc.add_paragraph(clean(line[4:]), style="Heading 2")
        elif line.startswith("| "):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                values = [part.strip() for part in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", part) for part in values):
                    rows.append(values)
                index += 1
            add_table(doc, rows)
            continue
        elif line.startswith("!["):
            match = re.match(r"!\[[^]]*\]\(([^)]+)\)", line)
            if not match:
                raise ValueError(line)
            path = ROOT / match.group(1)
            if not path.exists():
                raise FileNotFoundError(path)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(path), width=Inches(5.6))
            if pending_figure:
                caption = doc.add_paragraph(pending_figure, style="Caption")
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pending_figure = None
        elif re.match(r"^Figure \d+\. ", line):
            if line.startswith("Figure 1."):
                add_pipeline(doc)
                caption = doc.add_paragraph(line, style="Caption")
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                pending_figure = line
        elif line.startswith("SOURCE NASA TXT AND ZIP"):
            pass
        elif re.match(r"^Table \d+\. ", line):
            p = doc.add_paragraph(line, style="Caption")
            p.paragraph_format.keep_with_next = True
        else:
            p = doc.add_paragraph(clean(line))
            if cover:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if line.startswith("[") and re.match(r"^\[\d+\]", line):
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.first_line_indent = Inches(-0.25)
        index += 1

    doc.core_properties.title = "Predictive Remaining Useful Life Estimation of Aircraft Turbofan Engines"
    doc.core_properties.subject = "Machine Learning Group Activity Report draft"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
