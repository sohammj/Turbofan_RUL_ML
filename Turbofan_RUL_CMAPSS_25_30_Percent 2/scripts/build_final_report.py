"""Build the editable final Word report from the reviewed Markdown source."""

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "FINAL_PROJECT_REPORT.md"
OUTPUT = ROOT / "docs" / "FINAL_PROJECT_REPORT.docx"


def clean(text: str) -> str:
    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("**", "").replace("`", "").replace("*", "")
    return text.strip()


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    edges = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9D9D9")
        edges.append(edge)
    tc_pr.append(edges)


def add_table(doc, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.autofit = False
    widths = [1.03, 1.03, 1.03, 1.03, 1.03, 1.03, 1.03][:len(rows[0])]
    if len(widths) == 5:
        widths = [0.78, 0.92, 0.92, 1.27, 1.27]
    elif len(widths) == 7:
        widths = [0.65, 1.1, 0.88, 0.88, 1.1, 1.1, 0.88]
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].width = Inches(widths[index])
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[index].text = clean(value)
            borders(cells[index])
            if row_index == 0:
                shade(cells[index], "243B55")
                for run in cells[index].paragraphs[0].runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_index % 2 == 0:
                shade(cells[index], "F3F6F8")
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(28, 32, 36)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for style_name, size, before, after in (
        ("Title", 18, 0, 14),
        ("Heading 1", 13, 14, 6),
        ("Heading 2", 11, 10, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    title_style = doc.styles["Title"]
    title_ppr = title_style._element.get_or_add_pPr()
    title_borders = title_ppr.find(qn("w:pBdr"))
    if title_borders is not None:
        title_ppr.remove(title_borders)
    subtitle = doc.styles["Subtitle"]
    subtitle.font.name = "Arial"
    subtitle.font.size = Pt(10)
    subtitle.font.italic = False
    subtitle.font.color.rgb = RGBColor(0, 0, 0)
    caption_style = doc.styles["Caption"]
    caption_style.font.name = "Arial"
    caption_style.font.size = Pt(9)
    caption_style.font.bold = False
    caption_style.font.color.rgb = RGBColor(0, 0, 0)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("# "):
            doc.add_paragraph(clean(line[2:]), style="Title")
        elif line.startswith("## "):
            doc.add_paragraph(clean(line[3:]), style="Heading 1")
        elif line.startswith("> "):
            paragraph = doc.add_paragraph(clean(line[2:]))
            paragraph.style = "Subtitle"
            paragraph.paragraph_format.space_after = Pt(9)
        elif line.startswith("| "):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                parts = [part.strip() for part in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", part) for part in parts):
                    rows.append(parts)
                index += 1
            add_table(doc, rows)
            continue
        elif line.startswith("```"):
            index += 1
            commands = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                commands.append(lines[index].strip())
                index += 1
            for label, values in (
                ("Set up", commands[:3]),
                ("Run stages", commands[3:]),
            ):
                paragraph = doc.add_paragraph()
                paragraph.add_run(label + ": ").bold = True
                paragraph.add_run("  |  ".join(values))
                paragraph.paragraph_format.space_after = Pt(4)
        elif re.match(r"^\d+\. ", line):
            doc.add_paragraph(clean(line), style="Normal")
        else:
            paragraph = doc.add_paragraph(clean(line))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        index += 1

    doc.add_paragraph("Selected validation figure", style="Heading 1")
    figure = ROOT / "reports" / "robustness" / "repeated_validation_rmse.png"
    doc.add_picture(str(figure), width=Inches(6.35))
    caption = doc.add_paragraph(
        "Mean validation RMSE across three engine-wise splits. Error bars show the standard deviation across repetitions. This post-hoc audit did not change official-test model selection."
    )
    caption.style = "Caption"
    caption.paragraph_format.space_before = Pt(5)

    doc.core_properties.title = "Predictive Remaining Useful Life Estimation of Aircraft Turbofan Engines"
    doc.core_properties.subject = "Final academic machine learning project report"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
